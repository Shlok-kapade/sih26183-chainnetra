import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def c(hex_code):
    hex_code = hex_code.lstrip('#')
    return RGBColor(*(int(hex_code[i:i+2], 16) for i in (0, 2, 4)))

# Exact color palette from ppttemplate.pdf
NAVY_DARK   = c("0B132B")
NAVY_MED    = c("1C2541")
BLUE_PILL   = c("1D4E89")
BLUE_HEADER = c("1E3A8A")
BLUE_ACCENT = c("2563EB")
BLUE_LIGHT  = c("EFF6FF")
BLUE_CARD   = c("F0F4F8")
GREY_BG     = c("F8FAFC")
GREY_CARD   = c("F1F5F9")
GREY_BORDER = c("CBD5E1")
GREY_TEXT   = c("64748B")
DARK_TEXT   = c("0F172A")
WHITE       = c("FFFFFF")
ORANGE      = c("D97706")
ORANGE_DARK = c("EA580C")
ORANGE_BG   = c("FFFBEB")
ORANGE_CARD = c("FFEDD5")
GREEN_DARK  = c("15803D")
GREEN_BG    = c("DCFCE7")
GREEN_BORDER= c("86EFAC")
TEAL        = c("0F766E")
AMBER       = c("B45309")
PEACH_BG    = c("FFEDD5")

def add_shape(slide, shape_type, left, top, width, height, fill_color=None, line_color=None, line_width=Pt(1)):
    s = slide.shapes.add_shape(shape_type, left, top, width, height)
    s.line.fill.background()
    if fill_color:
        s.fill.solid()
        s.fill.fore_color.rgb = fill_color
    else:
        s.fill.background()
    if line_color:
        s.line.color.rgb = line_color
        s.line.width = line_width
    else:
        s.line.fill.background()
    return s

def add_box(slide, left, top, width, height, fill_color=None, line_color=None, line_width=Pt(0.75)):
    return add_shape(slide, MSO_SHAPE.RECTANGLE, left, top, width, height, fill_color, line_color, line_width)

def add_round_box(slide, left, top, width, height, fill_color=None, line_color=None, line_width=Pt(0.75)):
    return add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height, fill_color, line_color, line_width)

def add_pill(slide, text, left, top, width, height, bg_color=BLUE_HEADER, text_color=WHITE, font_size=8.5):
    add_round_box(slide, left, top, width, height, fill_color=bg_color)
    tb = slide.shapes.add_textbox(left, top - Inches(0.02), width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    r.font.bold = True
    r.font.size = Pt(font_size)
    r.font.color.rgb = text_color
    r.font.name = "Arial"

def add_textbox(slide, left, top, width, height, text="", font_size=9, bold=False, color=DARK_TEXT, align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    if text:
        r = p.add_run()
        r.text = text
        r.font.size = Pt(font_size)
        r.font.bold = bold
        r.font.color.rgb = color
        r.font.name = "Arial"
    return tb

def add_chevron(slide, left, top, width=Inches(0.12), height=Inches(0.2), color=GREY_BORDER):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "›"
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = color
    r.font.name = "Arial"

def clean_template_shapes(slide):
    for shape in list(slide.shapes):
        if shape.has_text_frame:
            txt = shape.text.strip()
            if any(k in txt for k in ["Detailed explanation", "Technologies to be used", "Analysis of the feasibility", "Potential impact", "Details / Links", "@SIH Idea"]):
                shape.left = Inches(-20)
            elif shape.name in ['TextBox 8', 'TextBox 9', 'Subtitle 3']:
                shape.left = Inches(-20)
            elif "Your\nTeam\nName" in shape.text or "Your Team Name" in shape.text or "Avengers" in shape.text or "code2trip" in shape.text:
                # Format the team oval beautifully like in ppttemplate.pdf
                shape.fill.solid()
                shape.fill.fore_color.rgb = c("1E293B")
                shape.line.fill.background()
                shape.text_frame.text = ""
                p = shape.text_frame.paragraphs[0]
                p.alignment = PP_ALIGN.CENTER
                r = p.add_run()
                r.text = "code2trip"
                r.font.size = Pt(11)
                r.font.bold = True
                r.font.color.rgb = WHITE
                r.font.name = "Arial"

def add_footer(slide, slide_num):
    # Cover any template artifact at bottom
    add_box(slide, Inches(0), Inches(6.9), Inches(13.333), Inches(0.6), fill_color=WHITE)
    # Dark blue bottom bar
    add_box(slide, Inches(0), Inches(7.15), Inches(13.333), Inches(0.35), fill_color=c("1E3A8A"))
    add_textbox(slide, Inches(0.5), Inches(7.20), Inches(10.0), Inches(0.25), 
                text=f"SIH 2026  |  PS SIH26183 — ChainNetra", font_size=8.5, bold=True, color=WHITE)
    add_textbox(slide, Inches(12.3), Inches(7.20), Inches(0.5), Inches(0.25), 
                text=str(slide_num), font_size=9, bold=True, color=WHITE, align=PP_ALIGN.RIGHT)

def build_slide1(prs):
    s = prs.slides[0]
    for shape in list(s.shapes):
        if shape.name in ['TextBox 9', 'Subtitle 3']:
            shape.left = Inches(-20)
        elif shape.has_text_frame and "Problem Statement ID" in shape.text:
            shape.left = Inches(-20)
    
    # Left column details
    details = [
        ("Problem Statement ID –", "SIH26183"),
        ("Problem Statement Title –", "Real-Time Identification of Fraud-Linked Cryptocurrency\nExchanges from Victim-Reported Suspect Wallet Addresses through Automated\nBlockchain Analytics"),
        ("Theme –", "Blockchain & Cybersecurity"),
        ("PS Category –", "Software"),
        ("Organisation –", "Ministry of Home Affairs"),
        ("Team ID –", "26183"),
        ("Team Name (Registered on portal) –", "code2trip"),
        ("Institute –", "code2trip Team")
    ]
    
    top = Inches(2.3)
    for label, val in details:
        tb = s.shapes.add_textbox(Inches(0.4), top, Inches(6.5), Inches(0.4))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        r1 = p.add_run()
        r1.text = label + "  "
        r1.font.bold = True
        r1.font.size = Pt(10)
        r1.font.color.rgb = BLUE_HEADER
        r1.font.name = "Arial"
        
        r2 = p.add_run()
        r2.text = val
        r2.font.bold = False
        r2.font.size = Pt(10)
        r2.font.color.rgb = DARK_TEXT
        r2.font.name = "Arial"
        
        if "\n" in val:
            top += Inches(0.55)
        else:
            top += Inches(0.28)
            
    # Bottom Left Solution Card
    card_left = Inches(0.4)
    card_top = Inches(5.6)
    card_w = Inches(6.2)
    card_h = Inches(1.3)
    add_round_box(s, card_left, card_top, card_w, card_h, fill_color=c("EBF3FA"))
    add_box(s, card_left, card_top, Inches(0.08), card_h, fill_color=c("1E3A8A"))
    
    tb = s.shapes.add_textbox(card_left + Inches(0.2), card_top + Inches(0.12), card_w - Inches(0.3), card_h - Inches(0.2))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    
    p0 = tf.paragraphs[0]
    r0 = p0.add_run()
    r0.text = "OUR SOLUTION\n"
    r0.font.size = Pt(8.5)
    r0.font.bold = True
    r0.font.color.rgb = BLUE_HEADER
    r0.font.name = "Arial"
    
    p1 = tf.add_paragraph()
    r1 = p1.add_run()
    r1.text = "ChainNetra\n"
    r1.font.size = Pt(22)
    r1.font.bold = True
    r1.font.color.rgb = NAVY_DARK
    r1.font.name = "Arial"
    
    p2 = tf.add_paragraph()
    r2 = p2.add_run()
    r2.text = "The address names no institution. The money does."
    r2.font.size = Pt(11)
    r2.font.color.rgb = GREY_TEXT
    r2.font.name = "Arial"

def build_slide2(prs):
    s = prs.slides[1]
    clean_template_shapes(s)
    
    # Title
    for shape in s.shapes:
        if shape.name == 'Title 1':
            shape.text_frame.text = ""
            p = shape.text_frame.paragraphs[0]
            r = p.add_run()
            r.text = "ChainNetra — Which Exchange Received the Money?"
            r.font.size = Pt(25)
            r.font.bold = True
            r.font.color.rgb = NAVY_DARK
            r.font.name = "Georgia"
            shape.top = Inches(0.55)
            shape.left = Inches(1.8)
            shape.width = Inches(9.5)
            
    # Subtitle bar
    sub_top = Inches(1.3)
    add_round_box(s, Inches(0.4), sub_top, Inches(12.5), Inches(0.42), fill_color=c("F8FAFC"), line_color=c("E2E8F0"))
    add_box(s, Inches(0.4), sub_top, Inches(0.08), Inches(0.42), fill_color=BLUE_HEADER)
    tb = s.shapes.add_textbox(Inches(0.6), sub_top + Inches(0.09), Inches(12.1), Inches(0.3))
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "A victim's wallet address goes in; the exchange that must receive the freeze request comes out, with the evidence attached."
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.color.rgb = NAVY_DARK
    r.font.name = "Arial"

    # Main Pipeline Section Header
    add_pill(s, "PROPOSED SOLUTION (DESCRIBE YOUR IDEA/SOLUTION/PROTOTYPE) · DETAILED EXPLANATION OF THE PROPOSED SOLUTION",
             Inches(0.4), Inches(1.85), Inches(9.2), Inches(0.28), bg_color=BLUE_PILL, font_size=8)
             
    # 6 Pipeline Step Cards
    steps = [
        ("01", "RETRIEVE", "We pull the transactions from public blockchain APIs. Every raw reply is hashed and timestamped as evidence."),
        ("02", "NORMALIZE", "TRON and Ethereum are turned into one common format — including the internal transfers most tracers miss."),
        ("03", "TRACE", "We follow the money hop by hop. Each address passes on only the victim's share of what left it, never more."),
        ("04", "GRAPH + PATTERNS", "Six detectors spot splitting, pooling, fast layering and peel chains."),
        ("05", "ATTRIBUTE", "We name the exchange behind a deposit address, with a tier, a confidence and the evidence for it."),
        ("06", "SCORE + REPORT", "19 risk signals, listed one by one, and a PDF case file with every transaction hash in it.")
    ]
    
    step_w = Inches(1.92)
    step_h = Inches(1.6)
    step_y = Inches(2.22)
    
    for i, (num, title, body) in enumerate(steps):
        x = Inches(0.4) + i * Inches(2.11)
        add_round_box(s, x, step_y, step_w, step_h, fill_color=c("EBF3FA"), line_color=c("DCE6F1"))
        add_box(s, x, step_y, Inches(0.06), step_h, fill_color=BLUE_HEADER)
        
        add_box(s, x + Inches(0.12), step_y + Inches(0.1), Inches(0.32), Inches(0.22), fill_color=BLUE_HEADER)
        add_textbox(s, x + Inches(0.12), step_y + Inches(0.12), Inches(0.32), Inches(0.22), text=num, font_size=8, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        
        add_textbox(s, x + Inches(0.5), step_y + Inches(0.13), step_w - Inches(0.55), Inches(0.22), text=title, font_size=8.5, bold=True, color=BLUE_HEADER)
        
        add_textbox(s, x + Inches(0.12), step_y + Inches(0.42), step_w - Inches(0.2), step_h - Inches(0.45), text=body, font_size=8, color=DARK_TEXT)
        
        if i < 5:
            add_chevron(s, x + step_w + Inches(0.04), step_y + Inches(0.65))

    # Bottom Sections
    bot_y = Inches(4.0)
    
    # Left Column: How it addresses the problem
    col1_x = Inches(0.4)
    col1_w = Inches(7.2)
    add_pill(s, "HOW IT ADDRESSES THE PROBLEM", col1_x, bot_y, Inches(3.5), Inches(0.28), bg_color=ORANGE_DARK, font_size=8)
    
    tbl_y = bot_y + Inches(0.35)
    
    add_box(s, col1_x, tbl_y, Inches(3.55), Inches(0.28), fill_color=c("F1F5F9"))
    add_textbox(s, col1_x + Inches(0.1), tbl_y + Inches(0.05), Inches(3.35), Inches(0.22), text="AN INVESTIGATOR TODAY", font_size=8.5, bold=True, color=GREY_TEXT)
    
    add_box(s, col1_x + Inches(3.65), tbl_y, Inches(3.55), Inches(0.28), fill_color=c("E0EDFB"))
    add_textbox(s, col1_x + Inches(3.75), tbl_y + Inches(0.05), Inches(3.35), Inches(0.22), text="WITH CHAINNETRA", font_size=8.5, bold=True, color=BLUE_HEADER)
    
    rows = [
        ("An address names nobody. A bank account number tells you the bank. A wallet address tells you nothing.",
         "Behaviour names it. An address that collects from many strangers and sweeps it onward is a deposit address."),
        ("Two or three hops, by hand. The money splits faster than a person can follow it.",
         "Every branch, automatically. The victim's share is carried correctly through each hop."),
        ("“Probably Binance.” A guess, with nothing behind it.",
         "A tier and its evidence. CONFIRMED, PROBABLE or UNATTRIBUTED — never rounded up to certainty."),
        ("6–12 hours, usually “unknown”. Long after the money has moved on. (field estimate)",
         "About 90 seconds on cached data. An evidenced answer to check, not a browser to click through.")
    ]
    
    ry = tbl_y + Inches(0.32)
    for left_txt, right_txt in rows:
        add_box(s, col1_x, ry, Inches(3.55), Inches(0.46), fill_color=c("F8FAFC"), line_color=c("E2E8F0"), line_width=Pt(0.5))
        add_textbox(s, col1_x + Inches(0.1), ry + Inches(0.04), Inches(3.35), Inches(0.4), text=left_txt, font_size=8, color=DARK_TEXT)
        
        add_box(s, col1_x + Inches(3.65), ry, Inches(3.55), Inches(0.46), fill_color=c("F0F7FF"), line_color=c("BAE6FD"), line_width=Pt(0.5))
        add_textbox(s, col1_x + Inches(3.75), ry + Inches(0.04), Inches(3.35), Inches(0.4), text=right_txt, font_size=8, color=DARK_TEXT)
        ry += Inches(0.49)
        
    leg_y = ry + Inches(0.05)
    legends = [
        (c("15803D"), "CONFIRMED in a published dataset"),
        (c("D97706"), "PROBABLE worked out, with a confidence"),
        (c("64748B"), "UNATTRIBUTED we do not know, and we say so")
    ]
    lx = col1_x
    for lcolor, ltext in legends:
        add_box(s, lx, leg_y + Inches(0.03), Inches(0.12), Inches(0.12), fill_color=lcolor)
        add_textbox(s, lx + Inches(0.18), leg_y, Inches(2.2), Inches(0.2), text=ltext, font_size=7.5, color=DARK_TEXT)
        lx += Inches(2.35)

    # Right Column: Innovation
    col2_x = Inches(7.8)
    col2_w = Inches(5.1)
    add_pill(s, "INNOVATION AND UNIQUENESS OF THE SOLUTION", col2_x, bot_y, col2_w, Inches(0.28), bg_color=GREEN_DARK, font_size=8)
    
    cards = [
        ("01", BLUE_HEADER, "THE THREE TIERS CAN NEVER BE MIXED UP", "The database itself rejects any claim that is not CONFIRMED, PROBABLE or UNATTRIBUTED."),
        ("02", ORANGE, "CONFIDENCE IS ONLY AS STRONG AS ITS WEAKEST LINK", "0.90 deposit × 0.80 exchange = 0.72. Never higher than that, and never above 0.95."),
        ("03", GREY_TEXT, "SEVEN NAMED REASONS A TRACE STOPS", "“We could not look further” is never reported as “the money stopped here”."),
        ("04", GREEN_DARK, "EVERY DETECTOR LISTS ITS OWN FALSE ALARMS", "A detector that cannot name an innocent explanation for the shape it flags is not allowed to run.")
    ]
    
    cy = bot_y + Inches(0.35)
    for num, num_color, title, desc in cards:
        add_round_box(s, col2_x, cy, col2_w, Inches(0.58), fill_color=c("F1F5F9"), line_color=c("E2E8F0"))
        add_round_box(s, col2_x + Inches(0.1), cy + Inches(0.1), Inches(0.35), Inches(0.35), fill_color=num_color)
        add_textbox(s, col2_x + Inches(0.1), cy + Inches(0.14), Inches(0.35), Inches(0.3), text=num, font_size=9, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        
        tb_c = s.shapes.add_textbox(col2_x + Inches(0.55), cy + Inches(0.06), col2_w - Inches(0.65), Inches(0.48))
        tf_c = tb_c.text_frame
        tf_c.word_wrap = True
        tf_c.margin_left = tf_c.margin_right = tf_c.margin_top = tf_c.margin_bottom = 0
        p1 = tf_c.paragraphs[0]
        r1 = p1.add_run()
        r1.text = title + "\n"
        r1.font.bold = True
        r1.font.size = Pt(8)
        r1.font.color.rgb = num_color
        
        p2 = tf_c.add_paragraph()
        r2 = p2.add_run()
        r2.text = desc
        r2.font.size = Pt(7.5)
        r2.font.color.rgb = DARK_TEXT
        
        cy += Inches(0.64)

    add_footer(s, 2)

def build_slide3(prs):
    s = prs.slides[2]
    clean_template_shapes(s)
    
    # Title
    for shape in s.shapes:
        if shape.name == 'Title 1':
            shape.text_frame.text = ""
            p = shape.text_frame.paragraphs[0]
            r = p.add_run()
            r.text = "TECHNICAL APPROACH"
            r.font.size = Pt(28)
            r.font.bold = True
            r.font.color.rgb = NAVY_DARK
            r.font.name = "Georgia"
            shape.top = Inches(0.6)
            shape.left = Inches(1.8)
            shape.width = Inches(9.5)

    # Left Section Header
    add_pill(s, "METHODOLOGY AND PROCESS FOR IMPLEMENTATION (FLOW CHARTS/IMAGES/ WORKING PROTOTYPE)",
             Inches(0.4), Inches(1.4), Inches(7.0), Inches(0.28), bg_color=BLUE_PILL, font_size=8)
             
    flow_cards = [
        ("INPUT", "The suspect address, as the victim reported it · the amount and the time, used to anchor the right transfer · the chain is worked out from the address itself"),
        ("RETRIEVE & NORMALIZE", "Rate-limited, retried, switched between providers · every raw reply hashed with SHA-256 and kept · TRON and Ethereum reduced to one Transfer model"),
        ("ANALYSE", "The trace runs in exact fractions with five stop conditions · the graph is built in memory and chokepoints found · six pattern detectors · deposit-address score"),
        ("OUTPUT", "The exchange named, with its tier and its evidence · 19 risk signals listed one by one · a PDF case file · anything incomplete comes back marked PARTIAL")
    ]
    
    fy = Inches(1.75)
    for i, (title, desc) in enumerate(flow_cards):
        add_round_box(s, Inches(0.4), fy, Inches(7.0), Inches(0.68), fill_color=c("EBF3FA"), line_color=c("DCE6F1"))
        add_box(s, Inches(0.4), fy, Inches(0.08), Inches(0.68), fill_color=BLUE_HEADER)
        
        add_textbox(s, Inches(0.6), fy + Inches(0.18), Inches(1.8), Inches(0.35), text=title, font_size=8.5, bold=True, color=BLUE_HEADER)
        add_textbox(s, Inches(2.4), fy + Inches(0.1), Inches(4.8), Inches(0.5), text=desc, font_size=8, color=DARK_TEXT)
        
        if i < 3:
            add_textbox(s, Inches(3.8), fy + Inches(0.65), Inches(0.3), Inches(0.22), text="▼", font_size=9, color=GREY_BORDER, align=PP_ALIGN.CENTER)
            
        fy += Inches(0.85)

    # Right Section: Technologies Used
    add_pill(s, "TECHNOLOGIES USED", Inches(7.6), Inches(1.4), Inches(5.3), Inches(0.28), bg_color=BLUE_PILL, font_size=8)
    
    techs = [
        ("Frontend", "React · TypeScript · Vite · Tailwind · Cytoscape.js"),
        ("Backend", "Python 3.12 · FastAPI · SQLAlchemy · Pydantic"),
        ("Data", "PostgreSQL 16 · Redis 7"),
        ("Chain data", "TronGrid · TronScan · Blockscout · Etherscan"),
        ("Analytics & AI", "NetworkX · LightGBM + SHAP (optional)"),
        ("Reports & tests", "ReportLab · pytest · Playwright")
    ]
    
    ty = Inches(1.75)
    for cat, val in techs:
        add_textbox(s, Inches(7.6), ty, Inches(1.4), Inches(0.2), text=cat, font_size=8.5, bold=True, color=BLUE_HEADER)
        add_textbox(s, Inches(9.0), ty, Inches(3.9), Inches(0.2), text=val, font_size=8.5, color=DARK_TEXT)
        ty += Inches(0.24)

    # Green Verified Card
    add_round_box(s, Inches(7.6), Inches(3.2), Inches(5.3), Inches(1.85), fill_color=GREEN_BG, line_color=GREEN_BORDER)
    add_box(s, Inches(7.6), Inches(3.2), Inches(0.08), Inches(1.85), fill_color=GREEN_DARK)
    
    tb_v = s.shapes.add_textbox(Inches(7.8), Inches(3.28), Inches(5.0), Inches(1.7))
    tf_v = tb_v.text_frame
    tf_v.word_wrap = True
    tf_v.margin_left = tf_v.margin_right = tf_v.margin_top = tf_v.margin_bottom = 0
    
    p0 = tf_v.paragraphs[0]
    r0 = p0.add_run()
    r0.text = "BUILT AND VERIFIED TODAY\n"
    r0.font.bold = True
    r0.font.size = Pt(9)
    r0.font.color.rgb = GREEN_DARK
    
    v_bullets = [
        "• Full pipeline runs: retrieval → tracing → graph → patterns → attribution → risk → PDF",
        "• 28 REST endpoints, and a workspace with seven tabs and an interactive fund-flow graph",
        "• 435 backend and 47 frontend tests pass with no network access at all",
        "• 94 source files, clean under ruff and mypy strict",
        "• Append-only audit log and evidence store, enforced by the database itself",
        "• Not claimed: the ML classifier is not built, and sign-in is single-factor"
    ]
    for b in v_bullets:
        pb = tf_v.add_paragraph()
        rb = pb.add_run()
        rb.text = b
        rb.font.size = Pt(7.5)
        rb.font.color.rgb = DARK_TEXT

    # Bottom Half Strip 1: How we spot a deposit address
    bot_y1 = Inches(5.2)
    add_round_box(s, Inches(0.4), bot_y1, Inches(12.5), Inches(0.65), fill_color=ORANGE_BG, line_color=c("FDE68A"))
    add_box(s, Inches(0.4), bot_y1, Inches(0.08), Inches(0.65), fill_color=ORANGE)
    
    tb_h = add_textbox(s, Inches(0.55), bot_y1 + Inches(0.12), Inches(3.2), Inches(0.45))
    tf_h = tb_h.text_frame
    tf_h.margin_left = tf_h.margin_right = tf_h.margin_top = tf_h.margin_bottom = 0
    p = tf_h.paragraphs[0]
    r = p.add_run()
    r.text = "HOW WE SPOT A DEPOSIT ADDRESS\n"
    r.font.bold = True
    r.font.size = Pt(8.5)
    r.font.color.rgb = ORANGE
    p2 = tf_h.add_paragraph()
    r2 = p2.add_run()
    r2.text = "Six weighted signals. 0.70 or more → PROBABLE."
    r2.font.size = Pt(7.5)
    r2.font.color.rgb = DARK_TEXT
    
    signals = [
        ("Sweeps to one place", "0.30"),
        ("Keeps almost nothing", "0.20"),
        ("Forwards within the hour", "0.15"),
        ("Three or more senders", "0.15"),
        ("One or two destinations", "0.10"),
        ("Never acts on its own", "0.10")
    ]
    sx = Inches(3.8)
    for stitle, sval in signals:
        add_round_box(s, sx, bot_y1 + Inches(0.08), Inches(1.4), Inches(0.48), fill_color=WHITE, line_color=c("FDE68A"))
        tb_s = add_textbox(s, sx + Inches(0.05), bot_y1 + Inches(0.12), Inches(1.3), Inches(0.4))
        tf_s = tb_s.text_frame
        tf_s.margin_left = tf_s.margin_right = tf_s.margin_top = tf_s.margin_bottom = 0
        p = tf_s.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = stitle + "\n"
        r.font.size = Pt(7)
        r.font.color.rgb = DARK_TEXT
        p2 = tf_s.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        r2 = p2.add_run()
        r2.text = sval
        r2.font.bold = True
        r2.font.size = Pt(8.5)
        r2.font.color.rgb = ORANGE
        sx += Inches(1.47)

    # Bottom Half Strip 2: Every number opens into its evidence
    bot_y2 = Inches(6.0)
    add_textbox(s, Inches(0.4), bot_y2 + Inches(0.15), Inches(3.0), Inches(0.4), 
                text="EVERY NUMBER OPENS INTO ITS EVIDENCE", font_size=8.5, bold=True, color=NAVY_DARK)
                
    chain_steps = [
        ("RAW RESPONSE", "stored, SHA-256"),
        ("TRANSFER", "one canonical row"),
        ("TRACE EDGE", "with its share"),
        ("ATTRIBUTION", "tier + evidence"),
        ("PDF APPENDIX", "hash + fetch time")
    ]
    cx = Inches(3.5)
    for i, (ctitle, cdesc) in enumerate(chain_steps):
        add_round_box(s, cx, bot_y2, Inches(1.6), Inches(0.55), fill_color=c("EBF3FA"), line_color=c("DCE6F1"))
        add_box(s, cx, bot_y2, Inches(0.06), Inches(0.55), fill_color=BLUE_HEADER)
        
        tb_c = add_textbox(s, cx + Inches(0.12), bot_y2 + Inches(0.08), Inches(1.4), Inches(0.4))
        tf_c = tb_c.text_frame
        tf_c.margin_left = tf_c.margin_right = tf_c.margin_top = tf_c.margin_bottom = 0
        p = tf_c.paragraphs[0]
        r = p.add_run()
        r.text = ctitle + "\n"
        r.font.bold = True
        r.font.size = Pt(7.5)
        r.font.color.rgb = BLUE_HEADER
        p2 = tf_c.add_paragraph()
        r2 = p2.add_run()
        r2.text = cdesc
        r2.font.size = Pt(7)
        r2.font.color.rgb = DARK_TEXT
        
        if i < 4:
            add_chevron(s, cx + Inches(1.62), bot_y2 + Inches(0.18))
        cx += Inches(1.85)

    add_footer(s, 3)

def build_slide4(prs):
    s = prs.slides[3]
    clean_template_shapes(s)
    
    # Title
    for shape in s.shapes:
        if shape.name == 'Title 1':
            shape.text_frame.text = ""
            p = shape.text_frame.paragraphs[0]
            r = p.add_run()
            r.text = "FEASIBILITY AND VIABILITY"
            r.font.size = Pt(28)
            r.font.bold = True
            r.font.color.rgb = NAVY_DARK
            r.font.name = "Georgia"
            shape.top = Inches(0.6)
            shape.left = Inches(1.8)
            shape.width = Inches(9.5)

    # Top Section: Analysis of the feasibility of the idea
    add_pill(s, "ANALYSIS OF THE FEASIBILITY OF THE IDEA", Inches(0.4), Inches(1.4), Inches(3.8), Inches(0.28), bg_color=GREEN_DARK, font_size=8)
    
    feas_cards = [
        ("IT ALREADY RUNS", "An address goes in; a named exchange, an itemised risk score and a PDF come out. 482 tests pass with no internet."),
        ("NO PAID DATA NEEDED", "Free public APIs and openly licensed datasets. No commercial forensics licence — which is why a district cyber cell could afford it."),
        ("ORDINARY DEPLOYMENT", "One server: Python, PostgreSQL, Redis and a web app. No cluster, no GPU and no graph database to run."),
        ("THE SAME ANSWER EVERY TIME", "Exact fractions, never floats. The same input gives the same answer — and every claim traces back to a stored, hashed API reply.")
    ]
    
    card_w = Inches(3.0)
    card_h = Inches(1.0)
    for i, (title, desc) in enumerate(feas_cards):
        x = Inches(0.4) + i * Inches(3.17)
        add_round_box(s, x, Inches(1.75), card_w, card_h, fill_color=GREEN_BG, line_color=GREEN_BORDER)
        add_box(s, x, Inches(1.75), Inches(0.08), card_h, fill_color=GREEN_DARK)
        
        tb = add_textbox(s, x + Inches(0.15), Inches(1.85), card_w - Inches(0.25), card_h - Inches(0.2))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = title + "\n"
        r.font.bold = True
        r.font.size = Pt(8.5)
        r.font.color.rgb = GREEN_DARK
        
        p2 = tf.add_paragraph()
        r2 = p2.add_run()
        r2.text = desc
        r2.font.size = Pt(7.5)
        r2.font.color.rgb = DARK_TEXT

    # Middle Section Headers
    mid_y = Inches(2.9)
    add_pill(s, "POTENTIAL CHALLENGES AND RISKS", Inches(0.4), mid_y, Inches(3.6), Inches(0.28), bg_color=ORANGE_DARK, font_size=8)
    add_pill(s, "STRATEGIES FOR OVERCOMING THESE CHALLENGES", Inches(6.6), mid_y, Inches(4.6), Inches(0.28), bg_color=BLUE_HEADER, font_size=8)
    
    risks_and_strategies = [
        ("TRON's free API is slow. We measured half a request per second per method, so a cold 200-address trace takes about 400 seconds — not 120.",
         "We work to an address budget, not to hope. About 60 new addresses inside 120 seconds; anything short of complete comes back marked PARTIAL. Cached and repeat traces stay fast."),
        ("Public TRON labels are few and dirty. Of 151 published exchange addresses we checked, 9 were not even valid TRON addresses.",
         "We check every label on the chain ourselves. 94 of the 151 survived that check. Each one records its source, its licence and the date we captured it."),
        ("A simple tracer cannot see internal transfers. On one sanctioned address, 173,600 ETH arrived where the ordinary transaction list showed nothing at all.",
         "We fetch them and treat them like any other transfer. Without that step a tracer misses 95.2% of the money that came in, and draws its biggest arrow the wrong way round."),
        ("A behaviour test raises false alarms. Payment processors, OTC desks and company sweep accounts all look like deposit addresses.",
         "Behaviour alone never gives CONFIRMED, and a detector that cannot name an innocent explanation for the shape it flags is not allowed to run."),
        ("A live demo depends on someone else's API. One rate-limit reply during a timed presentation ends the demo.",
         "Offline by default. Real chain data, captured and frozen, replayed from disk behind a visible “cached snapshot” banner. The whole suite runs with no network.")
    ]
    
    ry = mid_y + Inches(0.35)
    for risk, strat in risks_and_strategies:
        add_round_box(s, Inches(0.4), ry, Inches(6.0), Inches(0.62), fill_color=ORANGE_BG, line_color=c("FED7AA"))
        add_box(s, Inches(0.4), ry, Inches(0.06), Inches(0.62), fill_color=ORANGE_DARK)
        add_textbox(s, Inches(0.55), ry + Inches(0.06), Inches(5.75), Inches(0.5), text=risk, font_size=7.5, color=DARK_TEXT)
        
        add_round_box(s, Inches(6.6), ry, Inches(6.3), Inches(0.62), fill_color=c("F0F7FF"), line_color=c("BAE6FD"))
        add_box(s, Inches(6.6), ry, Inches(0.06), Inches(0.62), fill_color=BLUE_HEADER)
        add_textbox(s, Inches(6.75), ry + Inches(0.06), Inches(6.05), Inches(0.5), text=strat, font_size=7.5, color=DARK_TEXT)
        
        ry += Inches(0.67)

    add_textbox(s, Inches(0.4), Inches(6.75), Inches(12.0), Inches(0.25), 
                text="Every figure above is our own measurement, taken on 2026-09-05 and written up in the repository — not an estimate.",
                font_size=7.5, color=GREY_TEXT)

    add_footer(s, 4)

def build_slide5(prs):
    s = prs.slides[4]
    clean_template_shapes(s)
    
    # Title
    for shape in s.shapes:
        if shape.name == 'Title 1':
            shape.text_frame.text = ""
            p = shape.text_frame.paragraphs[0]
            r = p.add_run()
            r.text = "IMPACT AND BENEFITS"
            r.font.size = Pt(28)
            r.font.bold = True
            r.font.color.rgb = NAVY_DARK
            r.font.name = "Georgia"
            shape.top = Inches(0.6)
            shape.left = Inches(1.8)
            shape.width = Inches(9.5)

    # Top Section: Potential impact on the target audience
    add_pill(s, "POTENTIAL IMPACT ON THE TARGET AUDIENCE", Inches(0.4), Inches(1.4), Inches(3.8), Inches(0.28), bg_color=BLUE_PILL, font_size=8)
    
    stats = [
        ("22,68,346", "cybercrime incidents on NCRP in 2024, up 42% on 2023"),
        ("₹22,845.73 cr", "reported lost to cybercrime in 2024"),
        ("₹7,130 cr", "saved through CFCFRMS across 23.02 lakh complaints")
    ]
    
    stat_x = Inches(0.4)
    for num, lbl in stats:
        tb = add_textbox(s, stat_x, Inches(1.75), Inches(3.9), Inches(0.65))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = num + "\n"
        r.font.bold = True
        r.font.size = Pt(16)
        r.font.color.rgb = BLUE_HEADER
        
        p2 = tf.add_paragraph()
        r2 = p2.add_run()
        r2.text = lbl
        r2.font.size = Pt(8)
        r2.font.color.rgb = DARK_TEXT
        
        stat_x += Inches(4.2)
        
    add_textbox(s, Inches(8.5), Inches(2.4), Inches(4.4), Inches(0.2), 
                text="Ministry of Home Affairs, Lok Sabha Unstarred Question 432, answered 2 December 2025", 
                font_size=7, color=GREY_TEXT, align=PP_ALIGN.RIGHT)

    # 3 Persona Cards
    personas = [
        ("PRIMARY USER", "DISTRICT / STATE CYBER CELL INVESTIGATOR", "Not a blockchain expert, and does not need to become one. Types in the address the victim reported and gets back a named institution to serve, the hops in between, and the hashes to attach to the request."),
        ("SECONDARY", "I4C / CENTRAL ANALYST", "Sees the same deposit address turn up in complaints from different states — which turns twenty unconnected case files into one operation."),
        ("SECONDARY", "SUPERVISORY OFFICER", "Sorts a queue by where the money still seems to be, and can read why any finding was made without having to ask an engineer.")
    ]
    
    px = Inches(0.4)
    for ptype, ptitle, pbody in personas:
        add_round_box(s, px, Inches(2.65), Inches(3.95), Inches(1.15), fill_color=c("EBF3FA"), line_color=c("DCE6F1"))
        add_box(s, px, Inches(2.65), Inches(0.08), Inches(1.15), fill_color=BLUE_HEADER)
        
        tb = add_textbox(s, px + Inches(0.15), Inches(2.73), Inches(3.7), Inches(1.0))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = ptype + "\n"
        r.font.size = Pt(7)
        r.font.bold = True
        r.font.color.rgb = BLUE_HEADER
        
        p2 = tf.add_paragraph()
        r2 = p2.add_run()
        r2.text = ptitle + "\n"
        r2.font.bold = True
        r2.font.size = Pt(8.5)
        r2.font.color.rgb = NAVY_DARK
        
        p3 = tf.add_paragraph()
        r3 = p3.add_run()
        r3.text = pbody
        r3.font.size = Pt(7.5)
        r3.font.color.rgb = DARK_TEXT
        
        px += Inches(4.25)

    # Bottom Section Header
    add_pill(s, "BENEFITS OF THE SOLUTION (SOCIAL, ECONOMIC, ENVIRONMENTAL, ETC.)", 
             Inches(0.4), Inches(3.95), Inches(4.8), Inches(0.28), bg_color=GREEN_DARK, font_size=8)

    # Today vs With ChainNetra Strip
    comp_y = Inches(4.3)
    add_round_box(s, Inches(0.4), comp_y, Inches(5.8), Inches(0.85), fill_color=c("F8FAFC"), line_color=c("E2E8F0"))
    add_box(s, Inches(0.4), comp_y, Inches(0.08), Inches(0.85), fill_color=GREY_BORDER)
    tb_today = add_textbox(s, Inches(0.55), comp_y + Inches(0.08), Inches(5.5), Inches(0.7))
    tf_today = tb_today.text_frame
    tf_today.word_wrap = True
    p = tf_today.paragraphs[0]
    r = p.add_run()
    r.text = "HOW IT IS DONE TODAY\n"
    r.font.bold = True
    r.font.size = Pt(8)
    r.font.color.rgb = GREY_TEXT
    p2 = tf_today.add_paragraph()
    r2 = p2.add_run()
    r2.text = "A block explorer is opened and clicked through two or three hops by hand. Six to twelve hours per address, and it usually ends in “unknown”.\n"
    r2.font.size = Pt(7.5)
    r2.font.color.rgb = DARK_TEXT
    p3 = tf_today.add_paragraph()
    r3 = p3.add_run()
    r3.text = "Field estimate. No measured study exists, and we do not claim one."
    r3.font.size = Pt(7)
    r3.font.color.rgb = GREY_TEXT

    add_chevron(s, Inches(6.28), comp_y + Inches(0.35))

    add_round_box(s, Inches(6.5), comp_y, Inches(6.4), Inches(0.85), fill_color=c("F0F7FF"), line_color=c("BAE6FD"))
    add_box(s, Inches(6.5), comp_y, Inches(0.08), Inches(0.85), fill_color=BLUE_HEADER)
    tb_cn = add_textbox(s, Inches(6.65), comp_y + Inches(0.08), Inches(6.1), Inches(0.7))
    tf_cn = tb_cn.text_frame
    tf_cn.word_wrap = True
    p = tf_cn.paragraphs[0]
    r = p.add_run()
    r.text = "WITH CHAINNETRA\n"
    r.font.bold = True
    r.font.size = Pt(8)
    r.font.color.rgb = BLUE_HEADER
    p2 = tf_cn.add_paragraph()
    r2 = p2.add_run()
    r2.text = "The automated stages finish in about ninety seconds on cached or warm data. The investigator checks an evidenced answer instead of producing one.\n"
    r2.font.size = Pt(7.5)
    r2.font.color.rgb = DARK_TEXT
    p3 = tf_cn.add_paragraph()
    r3 = p3.add_run()
    r3.text = "Measured on captured chain data. A cold live trace is limited by provider rate limits, and says so on screen."
    r3.font.size = Pt(7)
    r3.font.color.rgb = GREY_TEXT

    # 4 Benefit Cards
    ben_y = Inches(5.25)
    benefits = [
        ("REACHES THE EXCHANGE IN TIME", "A freeze request needs someone to send it to. Naming that exchange early is the difference between getting the money back and not."),
        ("FINDS VICTIMS WHO NEVER COMPLAINED", "Tracing backwards from a scam wallet lists the other addresses that paid into it — people who never filed a complaint."),
        ("NO PER-SEAT LICENCE COST", "Commercial forensics platforms cost more than most state cyber cells can pay. This runs on free public data, on one server."),
        ("AN ANSWER THAT CAN BE CHECKED", "Every number opens into its signals, its transactions, and the stored raw reply with its hash and the time it was fetched.")
    ]
    bx = Inches(0.4)
    for btitle, bdesc in benefits:
        add_round_box(s, bx, ben_y, Inches(3.0), Inches(0.75), fill_color=WHITE, line_color=c("E2E8F0"))
        add_box(s, bx, ben_y, Inches(0.06), Inches(0.75), fill_color=BLUE_HEADER)
        
        tb = add_textbox(s, bx + Inches(0.12), ben_y + Inches(0.06), Inches(2.8), Inches(0.65))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = btitle + "\n"
        r.font.bold = True
        r.font.size = Pt(7.5)
        r.font.color.rgb = BLUE_HEADER
        p2 = tf.add_paragraph()
        r2 = p2.add_run()
        r2.text = bdesc
        r2.font.size = Pt(7)
        r2.font.color.rgb = DARK_TEXT
        
        bx += Inches(3.17)

    # Where this system stops banner
    stop_y = Inches(6.12)
    add_round_box(s, Inches(0.4), stop_y, Inches(12.5), Inches(0.65), fill_color=PEACH_BG, line_color=c("FED7AA"))
    add_box(s, Inches(0.4), stop_y, Inches(0.08), Inches(0.65), fill_color=ORANGE_DARK)
    
    tb_st = add_textbox(s, Inches(0.55), stop_y + Inches(0.12), Inches(2.2), Inches(0.45))
    tf_st = tb_st.text_frame
    p = tf_st.paragraphs[0]
    r = p.add_run()
    r.text = "WHERE THIS SYSTEM STOPS\n"
    r.font.bold = True
    r.font.size = Pt(8)
    r.font.color.rgb = ORANGE_DARK
    p2 = tf_st.add_paragraph()
    r2 = p2.add_run()
    r2.text = "Said by us, before anyone asks."
    r2.font.size = Pt(7)
    r2.font.color.rgb = GREY_TEXT
    
    limits = [
        "It does not name a person. Public blockchain data cannot do that, and we will not pretend otherwise.",
        "It does not decide that fraud happened. A risk score is an argument built from named signals — not a probability of guilt.",
        "It is not plugged into NCRP or CFCFRMS. We offer an integration-ready API and claim nothing more than that."
    ]
    lim_x = Inches(2.9)
    for lim in limits:
        add_textbox(s, lim_x, stop_y + Inches(0.12), Inches(3.0), Inches(0.45), text=lim, font_size=7, color=DARK_TEXT)
        lim_x += Inches(3.2)

    add_footer(s, 5)

def build_slide6(prs):
    s = prs.slides[5]
    clean_template_shapes(s)
    
    # Title
    for shape in s.shapes:
        if shape.name == 'Title 1':
            shape.text_frame.text = ""
            p = shape.text_frame.paragraphs[0]
            r = p.add_run()
            r.text = "RESEARCH AND REFERENCES"
            r.font.size = Pt(28)
            r.font.bold = True
            r.font.color.rgb = NAVY_DARK
            r.font.name = "Georgia"
            shape.top = Inches(0.6)
            shape.left = Inches(1.8)
            shape.width = Inches(9.5)

    # Header pill
    add_pill(s, "DETAILS / LINKS OF THE REFERENCE AND RESEARCH WORK", Inches(0.4), Inches(1.4), Inches(4.5), Inches(0.28), bg_color=BLUE_PILL, font_size=8)
    
    col_w = Inches(3.95)
    col_h = Inches(4.5)
    col_y = Inches(1.75)
    
    # Col 1: Government Mandate, Scale and Law
    add_round_box(s, Inches(0.4), col_y, col_w, col_h, fill_color=c("F8FAFC"), line_color=c("E2E8F0"))
    add_box(s, Inches(0.4), col_y, Inches(0.08), col_h, fill_color=BLUE_HEADER)
    
    tb1 = add_textbox(s, Inches(0.55), col_y + Inches(0.1), col_w - Inches(0.25), col_h - Inches(0.2))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    
    p = tf1.paragraphs[0]
    r = p.add_run()
    r.text = "GOVERNMENT MANDATE, SCALE AND LAW\n\n"
    r.font.bold = True
    r.font.size = Pt(8.5)
    r.font.color.rgb = BLUE_HEADER
    
    items1 = [
        ("MHA / I4C — SOP for NCRP–CFCFRMS, 2 January 2026\n",
         "Where funds reach an exchange not onboarded to CFCFRMS, the officer “may conduct a VDA forensic analysis … to ascertain that which VASP or VDA Exchange has control of the wallet.” Issuance confirmed in Rajya Sabha USQ 553, 4 Feb 2026.\n",
         "mha.gov.in/MHA1/Par2017/pdfs/par2026-pdfs/RS04022026/553.pdf\n\n"),
        ("MHA — Lok Sabha Unstarred Q. 432, 2 December 2025\n",
         "22,68,346 NCRP incidents in 2024 (+42.08%); ₹22,845.73 crore reported lost; ₹7,130 crore saved via CFCFRMS.\n",
         "mha.gov.in/MHA1/Par2017/pdfs/par2025-pdfs/LS02122025/432.pdf\n\n"),
        ("Standing Committee on Home Affairs — 254th Report, 20 Aug 2025\n",
         "Recommends training in “digital forensics, blockchain analysis and virtual asset tracing”.\n",
         "sansad.in › Committee_File › ReportFile › 254_2025_8_12.pdf\n\n"),
        ("Ministry of Finance — Gazette S.O. 1072(E), 7 March 2023\n",
         "Brings virtual digital asset services under the PMLA, 2002.\n",
         "egazette.gov.in/WriteReadData/2023/244184.pdf\n\n"),
        ("PIB / FIU-IND, 1 October 2025\n",
         "50 VDA service providers registered; non-compliance notices issued to 25 offshore exchanges.\n",
         "pib.gov.in/PressReleasePage.aspx?PRID=2173758")
    ]
    for h, b, u in items1:
        ph = tf1.add_paragraph()
        rh = ph.add_run()
        rh.text = h
        rh.font.bold = True
        rh.font.size = Pt(7.5)
        rh.font.color.rgb = DARK_TEXT
        
        pb = tf1.add_paragraph()
        rb = pb.add_run()
        rb.text = b
        rb.font.size = Pt(7)
        rb.font.color.rgb = DARK_TEXT
        
        pu = tf1.add_paragraph()
        ru = pu.add_run()
        ru.text = u
        ru.font.size = Pt(6.5)
        ru.font.color.rgb = BLUE_ACCENT

    # Col 2: Why TRON and USDT
    add_round_box(s, Inches(4.6), col_y, col_w, col_h, fill_color=c("F8FAFC"), line_color=c("E2E8F0"))
    add_box(s, Inches(4.6), col_y, Inches(0.08), col_h, fill_color=ORANGE)
    
    tb2 = add_textbox(s, Inches(4.75), col_y + Inches(0.1), col_w - Inches(0.25), col_h - Inches(0.2))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    
    p = tf2.paragraphs[0]
    r = p.add_run()
    r.text = "WHY TRON AND USDT\n\n"
    r.font.bold = True
    r.font.size = Pt(8.5)
    r.font.color.rgb = ORANGE
    
    items2 = [
        ("UNODC — Casinos, Money Laundering, Underground Banking and Transnational Organized Crime in East and Southeast Asia, Jan 2024\n",
         "“USDT on the TRON blockchain has become a preferred choice for regional cyberfraud operations and money launderers alike due to its stability and the ease, anonymity, and low fees of its transactions.” (p. 50) The same page reports an audit finding 17.07 billion USDT tied to underground exchanges and criminal activity between Sept 2022 and Sept 2023.\n",
         "unodc.org/roseap › Casino_Underground_Banking_Report_2024.pdf\n\n"),
        ("TRM Labs — 2025 Crypto Crime Report, 10 February 2025\n",
         "58% of illicit crypto volume in 2024 was on TRON, ahead of Ethereum (24%) and Bitcoin (12%). Illicit volume overall, not investment fraud alone.\n",
         "trmlabs.com/reports-and-whitepapers/2025-crypto-crime-report\n\n"),
        ("Enforcement Directorate — press release, 20 November 2025\n",
         "Proceeds of a ₹285-crore cyber-fraud network turned into USDT over exchange P2P using non-KYC accounts; ₹8.46 crore attached across 92 accounts. Names USDT; the chain is not stated.\n",
         "enforcementdirectorate.gov.in › Press Release-PAO-Cyber Fraud Case\n\n"),
        ("Enforcement Directorate — press release, 26 September 2025\n",
         "An investment scam promising 5–6% monthly returns collected funds through payment aggregators and USDT; ₹391 crore attached across 185 accounts.\n",
         "enforcementdirectorate.gov.in › Press Release-Navab Hassan QFX")
    ]
    for h, b, u in items2:
        ph = tf2.add_paragraph()
        rh = ph.add_run()
        rh.text = h
        rh.font.bold = True
        rh.font.size = Pt(7.5)
        rh.font.color.rgb = DARK_TEXT
        
        pb = tf2.add_paragraph()
        rb = pb.add_run()
        rb.text = b
        rb.font.size = Pt(7)
        rb.font.color.rgb = DARK_TEXT
        
        pu = tf2.add_paragraph()
        ru = pu.add_run()
        ru.text = u
        ru.font.size = Pt(6.5)
        ru.font.color.rgb = BLUE_ACCENT

    # Col 3: Data Sources and Our Own Measurements
    add_round_box(s, Inches(8.8), col_y, col_w, col_h, fill_color=c("F8FAFC"), line_color=c("E2E8F0"))
    add_box(s, Inches(8.8), col_y, Inches(0.08), col_h, fill_color=GREEN_DARK)
    
    tb3 = add_textbox(s, Inches(8.95), col_y + Inches(0.1), col_w - Inches(0.25), col_h - Inches(0.2))
    tf3 = tb3.text_frame
    tf3.word_wrap = True
    
    p = tf3.paragraphs[0]
    r = p.add_run()
    r.text = "DATA SOURCES AND OUR OWN MEASUREMENTS\n\n"
    r.font.bold = True
    r.font.size = Pt(8.5)
    r.font.color.rgb = GREEN_DARK
    
    items3 = [
        ("TronGrid · Blockscout · Etherscan\n",
         "Transaction data for TRON and Ethereum, used under their API terms. We never scrape explorer label databases.\n",
         "developers.tron.network/reference/rate-limits · docs.blockscout.com\n\n"),
        ("OFAC SDN sanctioned digital-currency addresses\n",
         "405 addresses ingested (281 TRON, 124 Ethereum), extracted with the MIT-licensed 0xB10C tool.\n",
         "github.com/0xB10C/ofac-sanctioned-digital-currency-addresses\n\n"),
        ("Dune spellbook — TRON exchange addresses\n",
         "A lead list only. Its BSL 1.1 licence does not cover our use, so we re-establish every address ourselves and never load theirs in.\n",
         "github.com/duneanalytics/spellbook\n\n"),
        ("Our measurements, 2026-09-05 (written up in the repository)\n",
         "TronGrid sustains half a request per second per method; Blockscout took 195 requests without throttling. 9 of 151 published TRON exchange labels are not valid addresses; 94 survived our checks.\n",
         "docs/research/OQ-01-provider-rate-limits.md · OQ-08-tron-label-coverage.md\n\n"),
        ("Our design record — 21 architecture decisions and a written limits list\n",
         "Why TRON and Ethereum first; why value is shared out in proportion; why there is no graph database; why the database enforces the three tiers; and what we will not claim.\n",
         "docs/DECISIONS.md · docs/LIMITATIONS.md")
    ]
    for h, b, u in items3:
        ph = tf3.add_paragraph()
        rh = ph.add_run()
        rh.text = h
        rh.font.bold = True
        rh.font.size = Pt(7.5)
        rh.font.color.rgb = DARK_TEXT
        
        pb = tf3.add_paragraph()
        rb = pb.add_run()
        rb.text = b
        rb.font.size = Pt(7)
        rb.font.color.rgb = DARK_TEXT
        
        pu = tf3.add_paragraph()
        ru = pu.add_run()
        ru.text = u
        ru.font.size = Pt(6.5)
        ru.font.color.rgb = BLUE_ACCENT

    # Bottom Banner: How we use these sources
    add_round_box(s, Inches(0.4), Inches(6.38), Inches(12.35), Inches(0.55), fill_color=ORANGE_BG, line_color=c("FDE68A"))
    add_box(s, Inches(0.4), Inches(6.38), Inches(0.08), Inches(0.55), fill_color=ORANGE)
    
    tb_bot = add_textbox(s, Inches(0.55), Inches(6.41), Inches(12.1), Inches(0.5))
    tf_bot = tb_bot.text_frame
    tf_bot.word_wrap = True
    p = tf_bot.paragraphs[0]
    r = p.add_run()
    r.text = "HOW WE USE THESE SOURCES\n"
    r.font.bold = True
    r.font.size = Pt(8)
    r.font.color.rgb = ORANGE
    
    p2 = tf_bot.add_paragraph()
    r2 = p2.add_run()
    r2.text = "We never scrape explorer label databases — their terms forbid it. A dataset whose licence does not allow our use is treated as a lead to check ourselves, never loaded in. Every label we ship records its source, its licence and the date we captured it."
    r2.font.size = Pt(7.5)
    r2.font.color.rgb = DARK_TEXT

    add_footer(s, 6)

def main():
    template_path = "docs/SIH_2026_PPT_Template.pptx"
    output_path = "docs/ChainNetra_SIH26183_code2trip.pptx"
    
    prs = Presentation(template_path)
    
    # Remove slide 7 if present
    if len(prs.slides) > 6:
        try:
            prs.slides._sldIdLst.remove(list(prs.slides._sldIdLst)[6])
        except Exception:
            pass
            
    print("Building Slide 1...")
    build_slide1(prs)
    print("Building Slide 2...")
    build_slide2(prs)
    print("Building Slide 3...")
    build_slide3(prs)
    print("Building Slide 4...")
    build_slide4(prs)
    print("Building Slide 5...")
    build_slide5(prs)
    print("Building Slide 6...")
    build_slide6(prs)
    
    prs.save(output_path)
    print(f"Saved exact high-fidelity presentation to {output_path}")

if __name__ == "__main__":
    main()
