import sys
from pptx import Presentation

def inspect_pptx(filepath):
    prs = Presentation(filepath)
    for i, slide in enumerate(prs.slides):
        print(f"\n--- Slide {i+1} ---")
        for j, shape in enumerate(slide.shapes):
            shape_type = shape.shape_type if hasattr(shape, 'shape_type') else 'Unknown'
            name = shape.name
            text = ""
            if shape.has_text_frame:
                text = shape.text[:50].replace('\n', ' ')
            print(f"  Shape {j}: type={shape_type}, name='{name}', text='{text}'")

if __name__ == "__main__":
    inspect_pptx(sys.argv[1])
